# -*- coding: utf-8 -*-
"""Convierte la guía Markdown del estudiante a DOCX."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from docx.shared import Inches, Pt, RGBColor


def set_run_font(run, name="Calibri", size=11, bold=False, italic=False, code=False):
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = "Consolas" if code else name
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.get_or_add_rFonts()
    rFonts.set(qn("w:ascii"), "Consolas" if code else name)
    rFonts.set(qn("w:hAnsi"), "Consolas" if code else name)
    if code:
        run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)


def add_horizontal_line(paragraph):
    p = paragraph._p
    pPr = p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "94A3B8")
    pBdr.append(bottom)
    pPr.append(pBdr)


def add_inline_runs(paragraph, text, base_size=11):
    # Soporta **negrita**, `codigo` y texto normal (mezclado)
    pattern = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`)")
    parts = pattern.split(text)
    for part in parts:
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, size=base_size, bold=True)
        elif part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, size=base_size - 1, code=True)
        else:
            run = paragraph.add_run(part)
            set_run_font(run, size=base_size)


def add_code_block(doc, code):
    for line in code.splitlines() or [""]:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.left_indent = Inches(0.15)
        # Fondo suave vía sombreado
        shd = OxmlElement("w:shd")
        shd.set(qn("w:fill"), "F1F5F9")
        shd.set(qn("w:val"), "clear")
        p._p.get_or_add_pPr().append(shd)
        run = p.add_run(line if line != "" else " ")
        set_run_font(run, size=9, code=True)


def add_table(doc, rows):
    if not rows:
        return
    cols = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=cols)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(cols):
            cell = table.rows[i].cells[j]
            cell.text = ""
            p = cell.paragraphs[0]
            value = row[j] if j < len(row) else ""
            # Limpia markup simple de markdown en celdas
            value = value.replace("**", "").replace("`", "")
            run = p.add_run(value)
            set_run_font(run, size=10, bold=(i == 0))
    doc.add_paragraph()


def parse_table_lines(lines):
    rows = []
    for line in lines:
        if re.match(r"^\s*\|?\s*:?-{3,}", line):
            continue
        cells = [c.strip() for c in line.strip().strip("|").split("|")]
        rows.append(cells)
    return rows


def convert(md_path: Path, docx_path: Path, titulo: str):
    text = md_path.read_text(encoding="utf-8")
    lines = text.splitlines()

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.9)
    section.bottom_margin = Inches(0.9)
    section.left_margin = Inches(0.95)
    section.right_margin = Inches(0.95)

    # Portada corta
    h = doc.add_paragraph()
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = h.add_run(titulo)
    set_run_font(run, size=22, bold=True)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = sub.add_run("Lenguaje de programación avanzado 2")
    set_run_font(run, size=12, italic=True)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = note.add_run(
        "Continúa después del documento «Creación de proyecto laravel»."
    )
    set_run_font(run, size=10)
    add_horizontal_line(doc.add_paragraph())

    i = 0
    # Saltar el H1 inicial del markdown (ya va en portada)
    if lines and lines[0].startswith("# "):
        i = 1
        while i < len(lines) and lines[i].strip() == "":
            i += 1

    in_code = False
    code_buf = []
    table_buf = []

    def flush_table():
        nonlocal table_buf
        if table_buf:
            add_table(doc, parse_table_lines(table_buf))
            table_buf = []

    while i < len(lines):
        line = lines[i]

        if line.startswith("```"):
            flush_table()
            if not in_code:
                in_code = True
                code_buf = []
            else:
                in_code = False
                add_code_block(doc, "\n".join(code_buf))
                code_buf = []
            i += 1
            continue

        if in_code:
            code_buf.append(line)
            i += 1
            continue

        if line.strip().startswith("|"):
            table_buf.append(line)
            i += 1
            continue
        else:
            flush_table()

        if line.strip() == "---":
            add_horizontal_line(doc.add_paragraph())
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_heading(line[2:].strip(), level=1)
            for run in p.runs:
                set_run_font(run, size=18, bold=True)
            i += 1
            continue

        if line.startswith("## "):
            p = doc.add_heading(line[3:].strip(), level=2)
            for run in p.runs:
                set_run_font(run, size=14, bold=True)
            i += 1
            continue

        if line.startswith("### "):
            p = doc.add_heading(line[4:].strip(), level=3)
            for run in p.runs:
                set_run_font(run, size=12, bold=True)
            i += 1
            continue

        if line.startswith("#### "):
            p = doc.add_heading(line[5:].strip(), level=4)
            for run in p.runs:
                set_run_font(run, size=11, bold=True)
            i += 1
            continue

        # Checklist / lista
        m_check = re.match(r"^- \[( |x|X)\] (.+)$", line)
        if m_check:
            checked = m_check.group(1).lower() == "x"
            prefix = "☑ " if checked else "☐ "
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.2)
            run = p.add_run(prefix)
            set_run_font(run, size=11)
            add_inline_runs(p, m_check.group(2))
            i += 1
            continue

        m_ul = re.match(r"^- (.+)$", line)
        if m_ul:
            p = doc.add_paragraph(style="List Bullet")
            # limpiar estilo default y poner runs
            p.clear()
            add_inline_runs(p, m_ul.group(1))
            i += 1
            continue

        m_ol = re.match(r"^(\d+)\. (.+)$", line)
        if m_ol:
            p = doc.add_paragraph(style="List Number")
            p.clear()
            add_inline_runs(p, m_ol.group(2))
            i += 1
            continue

        if line.strip() == "":
            i += 1
            continue

        p = doc.add_paragraph()
        add_inline_runs(p, line.strip())
        i += 1

    flush_table()
    if in_code and code_buf:
        add_code_block(doc, "\n".join(code_buf))

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))
    print(f"OK: {docx_path}")


def main():
    base = Path(__file__).resolve().parent
    materia = Path(
        r"d:\Enseñanza\Uniremintong\materias\2026\Lenguahe de programacion avanzado 2"
    )

    convert(
        base / "GUIA-ESTUDIANTE-PASO-A-PASO.md",
        materia / "GUIA-ESTUDIANTE-Cine-Paso-a-Paso.docx",
        "Laboratorio: sistema de cine (CRUD completo)",
    )
    convert(
        base / "GUIA-CLASE.md",
        materia / "GUIA-CLASE-explicar-cine.docx",
        "Guía de clase: CRUD del cine",
    )
    # Copia también dentro del proyecto
    convert(
        base / "GUIA-ESTUDIANTE-PASO-A-PASO.md",
        base / "GUIA-ESTUDIANTE-Cine-Paso-a-Paso.docx",
        "Laboratorio: sistema de cine (CRUD completo)",
    )


if __name__ == "__main__":
    main()
